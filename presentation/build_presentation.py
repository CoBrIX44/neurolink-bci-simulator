"""Build an editable 16:9 classroom deck using the local python-pptx skill.

Default brief: undergraduate Biology for Engineers; 16 slides, 12–15 minutes.
Run with the bundled document Python runtime; signal_examples.json is generated
with the project's Python 3.12 environment. No external generation API is used.
"""
from pathlib import Path
import json
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.chart.data import XyChartData, CategoryChartData

ROOT = Path(__file__).resolve().parent
DATA = json.loads((ROOT / 'signal_examples.json').read_text())
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BG, INK, TEAL, CORAL, MUTED, WHITE = 'F4F7F7', '16313B', '087E8B', 'D76B50', '526A72', 'FFFFFF'
team = [
 ('Aditya Dugar','25BRTSN003','Biology, motivation & introduction','1, 3–4'),
 ('Kushal Chaudhary','25BTRSN025','Acquisition concepts & synthetic signals','5–6'),
 ('Prince Jatheen Rout','25BTRSN034','Filtering, artifacts & scientific limits','7, 14'),
 ('Raunak Bothra','25BTRSN037','Features & machine-learning pipeline','8–9'),
 ('Vijaya Rupesh N','25BTRSN053','Validation, dashboard & live demo','10–12'),
 ('Yashas L Rao','25BTRSN055','Assistive-tech case study & learning review','13, 15–16')]
guide = ['# NeuroLink — Team 2 presenter guide\n', 'Target duration: 12–15 minutes plus questions. Roles below are assigned responsibilities for preparation and presentation, not a claim of individually verified past work. Rehearse and confirm ownership as a team.\n']

def text(s, value, x, y, w, h, size=22, color=INK, bold=False):
    shape = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(.03)
    tf.margin_top = tf.margin_bottom = Inches(.02)
    for i, line in enumerate(value.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.name, p.font.size = 'Arial', Pt(size)
        p.font.bold, p.font.color.rgb = bold, RGBColor.from_string(color)
        p.space_after = Pt(6)
    return shape

def rect(s,x,y,w,h,fill):
    shape=s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb=RGBColor.from_string(fill)
    shape.line.fill.background()
    return shape

def slide(title, owner, notes, subtitle=''):
    s=prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid(); s.background.fill.fore_color.rgb=RGBColor.from_string(BG)
    rect(s,0,0,.16,7.5,TEAL)
    text(s,'NEUROLINK  /  BIOLOGY FOR ENGINEERS',.55,.25,11,.3,12,TEAL,True)
    text(s,title,.55,.84,12.1,.75,32,INK,True)
    if subtitle: text(s,subtitle,.58,1.6,12,.62,17,MUTED)
    rect(s,.55,6.91,12.2,.015,'D2DFE1')
    text(s,'Team 2 · CE-SE-A | Synthetic signals only · Not a medical device',.58,7.06,9.6,.25,10,MUTED)
    text(s,f'{len(prs.slides):02d}',12.1,7.03,.6,.3,12,TEAL,True)
    s.notes_slide.notes_text_frame.text=f'Presenter: {owner}\n\n{notes}'
    guide.append(f'## {len(prs.slides)}. {title}\nPresenter: {owner}\n\n{notes}\n')
    return s

def card(s,x,y,w,title,body,accent=TEAL):
    rect(s,x,y,w,3.3,WHITE); rect(s,x,y,w,.07,accent)
    text(s,title,x+.2,y+.24,w-.4,.65,24,accent,True)
    text(s,body,x+.2,y+1.02,w-.4,2.05,21)

def banner(s,value,y=6.15):
    rect(s,.6,y,12.1,.5,INK); text(s,value,.8,y+.07,11.7,.36,16,WHITE)

def waveform(s, labels, raw=False, x=.65,y=2.25,w=8.0,h=3.7):
    d=XyChartData()
    for label in labels:
        row=DATA[label]
        for field in (['raw','filtered'] if raw else ['raw']):
            series=d.add_series(f'{label}: {field}')
            for t,v in zip(row['time'],row[field]): series.add_data_point(t,v)
    chart=s.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER_LINES_NO_MARKERS,Inches(x),Inches(y),Inches(w),Inches(h),d).chart
    chart.has_legend=True; chart.legend.position=XL_LEGEND_POSITION.TOP
    chart.legend.font.size=Pt(12); chart.legend.include_in_layout=True
    for axis, title in [(chart.category_axis,'Time (seconds)'),(chart.value_axis,'Amplitude (arbitrary units)')]:
        axis.has_title=True; axis.axis_title.text_frame.text=title
        axis.axis_title.text_frame.paragraphs[0].font.size=Pt(13)
        axis.tick_labels.font.size=Pt(11)
    for i,series in enumerate(chart.series):
        series.format.line.color.rgb=RGBColor.from_string([MUTED,TEAL,CORAL][i%3])
        series.format.line.width=Pt(1.3)
    return chart

s=slide('NeuroLink', 'Aditya Dugar',
 'Introduce Team 2 and the Biology for Engineers experiential-learning project. Explain that the goal is to learn the software stages of a BCI without a headset. Say clearly that all displayed signals are generated and that left/right are artificial labels. The presentation moves from biology to the processing pipeline, a demonstration, and the limits of what the result proves.',
 'An educational brain–computer interface simulator')
rect(s,.6,2.4,7.5,3.72,INK)
text(s,'From a noisy waveform\nto an artificial command',.9,2.77,6.9,1.7,35,WHITE,True)
text(s,'Generate → Filter → Extract → Classify → Visualize',.93,5.04,6.8,.7,20,'80D7DA')
text(s,'EXPERIENTIAL LEARNING',8.6,2.8,4,.4,17,TEAL,True)
text(s,'Team No. 2\nClass: CE-SE-A\nBiology for Engineers',8.6,3.45,4,1.9,25)
banner(s,'Educational prototype using simulated EEG-like signals — not live brain data.')

s=slide('Team 2 · contribution & presentation plan','All members',
 'These are assigned responsibilities, not a historical authorship claim. Each member should understand their module, rehearse the explanation, and review the associated code or evidence. Aditya frames the biology; Kushal explains signals; Prince covers filtering and limitations; Raunak covers features and learning; Vijaya demonstrates and validates; Yashas connects the case study and learning outcomes. Use this table to organize rehearsal.',
 'Assigned responsibilities for preparation, explanation and demonstration')
cols=[.7,3.2,5.5,11.3]
for x,v,w in zip(cols,['MEMBER','STUDENT ID','RESPONSIBILITY','SLIDES'],[2.4,2.2,5.6,1.2]): text(s,v,x,2.24,w,.35,12,TEAL,True)
for i,(name,id,role,nums) in enumerate(team):
    y=2.83+i*.56
    if i%2==0: rect(s,.62,y-.04,12.05,.51,WHITE)
    for x,v,w in zip(cols,[name,id,role,nums],[2.4,2.2,5.6,1.3]): text(s,v,x,y,w,.5,16)
text(s,'Shared: integration review, rehearsal and questions. Roles are a proposed allocation, not verified individual authorship.',.7,6.34,12,.4,13,MUTED)

s=slide('The learning challenge','Aditya Dugar',
 'Our learning question is how a noisy time-domain signal becomes a compact representation and a command. Real acquisition would require equipment and experimental controls beyond this prototype. A simulator makes the stages observable and repeatable. Explain success in terms of an understandable pipeline, reproducibility and honest boundaries, not a clinical performance target.',
 'How does a waveform become a usable command?')
card(s,.65,2.5,3.85,'Observe','Compare raw and filtered signals using the same time axis.')
card(s,4.75,2.5,3.85,'Measure','Extract two named frequency-band features from each example.')
card(s,8.85,2.5,3.85,'Interpret','Classify generated patterns and explain what the output can—and cannot—mean.',CORAL)
banner(s,'Learning objective: explain each transformation, then test its behavior.')

s=slide('Biology first: what a BCI measures','Aditya Dugar',
 'A real BCI measures neural activity and maps selected patterns to an output. EEG measures electrical potential differences through electrodes; it is affected by the recording setup and unwanted signal contributions. Our app does not perform that acquisition step. Make the distinction before showing any waveform so the audience does not mistake synthetic sinusoids for recordings. Hand over to Kushal to explain how the simulation replaces acquisition.',
 'Neural activity → measurement → signal processing → external output')
card(s,.7,2.45,5.85,'Real EEG pathway','Electrodes record voltage differences.\nSignals vary across people and sessions.\nEye, muscle and motion activity can affect recordings.')
card(s,6.8,2.45,5.85,'Our educational substitution','Python generates a waveform.\nAmplitude uses arbitrary units.\nNo headset, human participant or recorded EEG is involved.',CORAL)

s=slide('The complete software pipeline','Kushal Chaudhary',
 'Walk from left to right through the pipeline. Every example uses 512 samples. Filtering returns the same number of samples; feature extraction reduces that waveform to two numbers. The trained classifier maps those numbers to an artificial category. The Streamlit dashboard shows the intermediate evidence rather than hiding it behind a final label. Name NumPy, SciPy and scikit-learn as implementation tools, without treating a software library as biological validation.')
steps=[('01','Generate','512 samples'),('02','Filter','1–40 Hz'),('03','Extract','2 features'),('04','Classify','left / right'),('05','Display','plots + cursor')]
for i,(n,t,b) in enumerate(steps):
    x=.7+i*2.52
    rect(s,x,2.65,2.32,2.65,WHITE)
    text(s,n,x+.2,2.85,1.9,.55,32,TEAL,True)
    text(s,t,x+.2,3.65,1.95,.55,24,INK,True)
    text(s,b,x+.2,4.44,1.95,.45,17,MUTED)
text(s,'NumPy  ·  SciPy  ·  scikit-learn  ·  joblib  ·  Streamlit  ·  Plotly',.8,5.55,11.8,.45,19)
banner(s,'Python 3.12 · Local-first application with an optional Streamlit-hosted demo')

s=slide('Create two deliberately distinct patterns','Kushal Chaudhary',
 'These native charts contain actual outputs from generate_eeg using seed 7, not stock EEG images. Left emphasizes 8 Hz and right 22 Hz. The generator adds a weaker alternate component, 0.3 Hz drift, 60 Hz interference and Gaussian noise with randomized phase and amplitude. At 256 samples per second for two seconds, the output has 512 samples. Explain that the deliberately different patterns make classification easy; a fixed seed makes the example repeatable.',
 'Actual generated examples · seed 7 · arbitrary amplitudes')
waveform(s,['left','right'])
text(s,'256 Hz\n2 seconds\n512 samples',9.08,2.65,3.45,2.1,29,TEAL,True)
text(s,'left: emphasizes 8 Hz\nright: emphasizes 22 Hz\nNoise + drift + interference',9.08,4.85,3.5,1.1,18)
banner(s,'Left and right describe generator settings, not a person’s intention.')

s=slide('Filter the signal, not every artifact','Prince Jatheen Rout',
 'The chart compares the actual raw and filtered left example at seed 7. Explain the fourth-order Butterworth band-pass, represented as second-order sections. The algorithm runs forward and backward to produce zero-phase output. This uses future samples and therefore is an offline demonstration. Filtering attenuates out-of-band content but cannot distinguish an artifact inside the useful band from desired content. Edge effects are also possible in this short window.',
 'Butterworth band-pass · second-order sections · offline zero phase')
waveform(s,['left'],raw=True)
text(s,'Pass band\n1–40 Hz',9.07,2.62,3.4,1.25,30,TEAL,True)
text(s,'Same sample count\nFinite output checked\nFuture samples required\nEdge effects remain',9.07,4.1,3.4,1.85,20)
banner(s,'Motion, eye and muscle artifacts may overlap the passed frequencies.')

s=slide('Reduce 512 samples to two features','Raunak Bothra',
 'First subtract the mean, then calculate the real FFT. For each positive-frequency bin, square the magnitude normalized by the sample count. Average those values in each band. We choose a nonoverlapping convention: 1 inclusive to 12 exclusive, then 12 through 40 inclusive. These are mean powers per FFT bin, not integrated power or a PSD. The chart uses actual processed seed-7 examples. Explain that the relative power distribution gives the classifier useful separation.',
 'Center → real FFT → squared normalized magnitude → band averages')
d=CategoryChartData(); d.categories=['1 ≤ f < 12 Hz','12 ≤ f ≤ 40 Hz']
for label in ['left','right']: d.add_series(label, list(DATA[label]['features'].values()))
c=s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(.7), Inches(2.35), Inches(7.25), Inches(3.65),d).chart
c.has_legend=True; c.legend.position=XL_LEGEND_POSITION.TOP; c.legend.include_in_layout=True
c.value_axis.has_title=True; c.value_axis.axis_title.text_frame.text='Mean power (arbitrary units²)'
c.category_axis.tick_labels.font.size=Pt(14)
for ser,col in zip(c.series,[TEAL,CORAL]): ser.format.fill.solid(); ser.format.fill.fore_color.rgb=RGBColor.from_string(col)
text(s,'Stable feature names',8.3,2.5,4.2,.55,24,TEAL,True)
text(s,'mean_power_1_12_hz\nmean_power_12_40_hz',8.3,3.25,4.2,1.05,20)
text(s,'Mean |rFFT(x − mean(x)) / N|²\nwithin each band',8.3,4.65,4.2,1.0,19)
banner(s,'The 12 Hz bin is counted once. Empty bands raise an explicit error.')

s=slide('Train a reproducible classifier','Raunak Bothra',
 'The dataset has 200 examples in each class. A fixed-seed stratified split preserves balance: 300 examples for training and 100 for testing. StandardScaler is fitted on the training split only, preventing test statistics from entering training. Logistic regression learns the separation between the two synthetic feature patterns. A local joblib artifact can be loaded later; on a fresh hosted deployment the dashboard can build a cached in-memory synthetic model. Probabilities remain model outputs, not confidence about human intent.',
 'Balanced synthetic data · fixed random state · held-out evaluation')
card(s,.7,2.42,3.85,'400 examples','200 left + 200 right\nIndependent example seeds\nOne generator family')
card(s,4.76,2.42,3.85,'300 / 100 split','Stratified 75% / 25%\nFit scaler on training only\nHold test examples aside')
card(s,8.82,2.42,3.85,'Model pipeline','StandardScaler\nLogisticRegression\nSave/load with joblib')
banner(s,'Reproducibility supports teaching; it does not establish real EEG generalization.')

s=slide('Validation: what the results demonstrate','Vijaya Rupesh N',
 'The saved seed-42 training report contains 100 test examples: 50 per class, all correctly classified. Report the 100% figure only with the synthetic-only qualification. These classes were designed to be different, so this is not comparable to a clinical BCI result. Automated tests cover parameter validation, signal determinism, filter behavior, FFT band boundaries, dataset balance, saved models, dashboard interactions, model fallback and cursor bounds. Cite the actual local test run reported in the presenter guide.',
 'Source: saved synthetic metrics and locally executed automated tests')
text(s,'100 / 100',.85,2.5,5.4,1.05,58,TEAL,True)
text(s,'held-out synthetic examples correct',.9,3.65,5.2,.9,23)
text(s,'Confusion matrix\nActual → predicted',7.05,2.45,5,.9,23,INK,True)
for i,row in enumerate([['','left','right'],['left','50','0'],['right','0','50']]):
    for j,value in enumerate(row):
        x,y=7.1+j*1.55,3.55+i*.64
        rect(s,x,y,1.43,.55,TEAL if i==j and i>0 else WHITE)
        text(s,value,x+.14,y+.06,1.15,.44,21,WHITE if i==j and i>0 else INK)
text(s,'58 tests passed locally*',.9,5.25,5.7,.62,25,INK,True)
banner(s,'Synthetic accuracy is NOT evidence of real EEG decoding or clinical effectiveness.')
text(s,'*See presenter guide for the verification date and command.',.9,6.72,10,.18,9,MUTED)

s=slide('An interactive dashboard makes the stages visible','Vijaya Rupesh N',
 'This slide is a schematic of the implemented dashboard, not a screenshot. The sidebar selects the artificial class, reproducible seed and optional cursor. The output contains a Plotly raw/filtered chart, named feature values, prediction and class scores. The displayed result retains its original seed and label until processed again. Missing local model files trigger cached in-memory synthetic training; other failures are shown as actionable errors. The synthetic-data notice is always included in the app.',
 'Interface map · representative values from the actual seed-7 left example')
rect(s,.7,2.4,3,3.55,INK)
text(s,'INPUT',.95,2.65,2.5,.5,19,'80D7DA',True)
text(s,'Artificial class: left\nSeed: 7\nGenerate and process\nOptional simulated cursor',.95,3.36,2.55,2.15,20,WHITE)
rect(s,4,2.4,8.65,3.55,WHITE)
text(s,'Raw + filtered waveform comparison',4.3,2.65,7.95,.55,24,TEAL,True)
text(s,f"Low-band feature: {DATA['left']['features']['mean_power_1_12_hz']:.4f}\nHigh-band feature: {DATA['left']['features']['mean_power_12_40_hz']:.4f}",4.3,3.55,4.65,1.15,22)
text(s,f"Prediction: left\nModel score: {DATA['left']['scores']['left']:.4f}",9.05,3.55,3.3,1.15,22)
text(s,'Zoomable chart · explicit errors · scientific explanation',4.3,5.12,7.9,.6,20,MUTED)
banner(s,'Interface schematic; model scores do not measure actual human intention.')

s=slide('Live demonstration · about 90 seconds','Vijaya Rupesh N',
 'Before presenting, open the local dashboard or the working Streamlit deployment. Start at seed 7 with the cursor disabled. Generate left and point out the slower dominant oscillation and larger low-band feature. Generate right and compare the faster oscillation and higher upper-band feature. Enable the cursor and reset to 50: process left once for 40, then right once for 50. Changing controls alone should not move the cursor. If the network is unavailable, run locally using the command on this slide; the preceding charts provide a backup.',
 'Show the causal chain rather than only the final label')
for i,(a,b) in enumerate([('Choose left · seed 7','Compare raw/filtered traces and inspect the low-band feature.'),('Choose right · seed 7','Compare the feature balance and predicted artificial label.'),('Enable cursor · reset to 50','Left → 40; then right → 50. Bounds are 0–100.')]):
    y=2.35+i*1.08
    text(s,f'0{i+1}',.8,y,.85,.7,32,TEAL,True)
    text(s,a,1.8,y,10.5,.5,24,INK,True); text(s,b,1.8,y+.52,10.5,.5,19,MUTED)
banner(s,'Local backup: python -m streamlit run app.py')

s=slide('Real assistive technology: the BrainGate study','Yashas L Rao',
 'Keep this case study clearly separate from NeuroLink. Hochberg and colleagues reported in 2012 that two participants with tetraplegia controlled robotic reaching and grasping using implanted motor-cortex electrode arrays. This was an investigational intracortical neural-interface study, not scalp EEG. It illustrates the motivation for assistive BCI work but does not validate our classifier or simulator. Do not generalize a two-participant research demonstration to universal effectiveness. Source: https://pmc.ncbi.nlm.nih.gov/articles/PMC3640850/',
 'A separate published research example — not a capability of our simulator')
card(s,.7,2.4,5.85,'Published research','Two participants with tetraplegia\nImplanted motor-cortex electrodes\nRobotic reaching and grasping')
card(s,6.8,2.4,5.85,'NeuroLink','Generated sinusoidal patterns\nTwo FFT features and an artificial label\nOn-screen cursor simulation',CORAL)
text(s,'Hochberg et al., Nature 485, 372–375 (2012) · doi:10.1038/nature11076',.8,6.12,12,.45,16,TEAL)

s=slide('Scientific boundaries we must keep visible','Prince Jatheen Rout',
 'Use this slide to prevent overclaiming during questions. We have no measured neural signal, no verified motor-imagery label and no clinical validation. Filtering is not complete artifact removal and zero-phase filtering is not streaming. Synthetic train and test samples come from the same generator family. A future real-EEG project would require a documented licensed dataset, careful leakage controls such as subject/session separation as appropriate, and evaluation on genuinely held-out data. It would be a separate study.',
 'What works in this simulator does not establish what works on a person')
card(s,.7,2.45,3.85,'Data boundary','Synthetic only\nArbitrary amplitudes\nArtificial labels')
card(s,4.76,2.45,3.85,'Method boundary','Offline zero phase\nIncomplete artifact removal\nPedagogical band choices')
card(s,8.82,2.45,3.85,'Evidence boundary','Same generator family\nNo human-intent confidence\nNo clinical validation',CORAL)
banner(s,'Real EEG is future work requiring a separate dataset, protocol and evaluation.')

s=slide('Experiential learning: build, observe, explain','Yashas L Rao',
 'Frame these as learning takeaways and a reflection plan rather than invented student testimonials. Each member should be able to explain a transformation using evidence from the app. Ask the team what changes when seed changes, why a filter cannot remove every artifact, and why a high test score can be misleading. Close by restating that the outcome is a working and testable educational signal-processing pipeline. Invite questions on either the engineering implementation or the distinction from real BCI research.',
 'Connect biology concepts to engineering decisions')
card(s,.7,2.43,3.85,'Build & inspect','Modular signal processing\nReproducible experiments\nVisible intermediate outputs')
card(s,4.76,2.43,3.85,'Test & question','Check shape and finite values\nTest frequency behavior\nChallenge accuracy claims')
card(s,8.82,2.43,3.85,'Explain & reflect','Identify artifacts\nExplain synthetic data\nState scientific limits')
banner(s,'Discussion: why can 100% synthetic accuracy still tell us very little about real EEG?')

s=slide('References & project resources','Yashas L Rao / all members',
 'Use the repository and local documentation for implementation-specific statements, and the cited primary research for the assistive case study. The class should not treat framework documentation as evidence of biological efficacy. The app deployment was reported successful by the team during setup; the deck does not claim an independent cloud availability test. Confirm the current hosted link from the repository About section before the presentation. Thank the audience and invite questions.')
refs=[('Primary research','Hochberg et al. (2012). Reach and grasp by people with tetraplegia using a neurally controlled robotic arm. Nature 485, 372–375.','https://pmc.ncbi.nlm.nih.gov/articles/PMC3640850/'),('Project repository','Source code, README, specifications and setup instructions.','https://github.com/CoBrIX44/neurolink-bci-simulator'),('Methods & evidence','PRD.md · docs/SCIENCE_AND_USER_GUIDE.md · docs/IMPLEMENTATION_NOTES.md · models/neurolink.metrics.json',''),('Software documentation','SciPy signal filtering · scikit-learn pipelines · Streamlit application framework','https://docs.scipy.org/doc/scipy/reference/signal.html')]
for i,(a,b,url) in enumerate(refs):
    y=2.2+i*1.03
    text(s,a,.75,y,2.5,.5,19,TEAL,True)
    text(s,b,3.25,y,9.25,.65,17)
    if url:
        sh=text(s,url,3.25,y+.61,9.25,.3,12,TEAL)
        sh.text_frame.paragraphs[0].runs[0].hyperlink.address=url

out=ROOT/'NeuroLink_Team_2_Experiential_Learning.pptx'
prs.save(out)
guide.append('## Verification record\n12 October 2026: `.\\.venv\\Scripts\\python.exe -m pytest -q` returned 58 passed in 3.98 seconds. Charts use actual seed-7 generated examples. Evaluation uses the saved seed-42 model report: 100 held-out synthetic examples, all correctly classified. This is not real EEG performance.\n')
(ROOT/'PRESENTER_GUIDE.md').write_text('\n'.join(guide),encoding='utf-8')
print(f'Saved {len(prs.slides)} slides: {out}')


